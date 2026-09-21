#!/usr/bin/env python3
import argparse, json, re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FIELDS=[
("initiative_name","Initiative name","Untitled initiative"),("one_sentence_promise","One-sentence promise","TBD"),("portfolio","Portfolio / business unit","TBD"),("executive_audience","Executive audience and decision owner","TBD"),("decision_requested","Decision requested","Approve and prioritize"),
("primary_customer","Primary customer or user","TBD"),("customer_problem","Customer problem","TBD"),("current_workflow","Current workflow and friction","TBD"),("problem_evidence","Evidence or hypothesis","TBD"),("customer_outcome","Measurable customer outcome","TBD"),("behavior_change","Today to future behavior change","TBD"),("value_proposition","Shortest defensible value proposition","TBD"),("north_star","North-star metric","TBD"),
("solution","Proposed experience or capability","TBD"),("availability_now","Available now","TBD"),("near_term","Near term","TBD"),("medium_term","Medium term","TBD"),("beyond","Beyond","TBD"),("right_to_win","Right to win / differentiation","TBD"),("out_of_scope","Out of scope","TBD"),("portfolio_leverage","Portfolio leverage","TBD"),("dependencies","Dependencies and commercial gates","TBD"),("operating_model","Cross-functional operating model","TBD"),
("growth_mechanism","Primary and secondary growth mechanisms","TBD"),("buying_committee","Buyer, champions, users, and approvers","TBD"),("pricing_hypothesis","Pricing and packaging hypothesis","TBD"),("financial_scenarios","Supported financial scenario inputs","TBD"),("phase_gates","Phase gates and evidence","TBD"),("pilot_plan","Design-partner or pilot plan","TBD"),
("risks","Top risks and mitigations","TBD"),("governance","Trust, security, compliance, and approval controls","TBD"),("recommendation","Final recommendation","APPROVE AND PRIORITIZE, subject to evidence gates."),("validate_next","Assumptions and what to validate next","TBD"),("status","Status and planning horizon","Business initiative - executive proposal")]

def ask():
    print("Business Initiative interview. Enter TBD for unknown items.\n")
    d={}
    groups=[("Identity",FIELDS[:5]),("Customer and outcome",FIELDS[5:13]),("Solution and roadmap",FIELDS[13:23]),("Commercial proof",FIELDS[23:29]),("Risk and decision",FIELDS[29:])]
    for title,items in groups:
        print(f"\n== {title} ==")
        for key,label,default in items:
            v=input(f"{label} [{default}]: ").strip()
            d[key]=v or default
    return d

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_cell_text(cell,text,bold=False,color=None):
    cell.text=''; p=cell.paragraphs[0]; r=p.add_run(str(text)); r.bold=bold
    if color: r.font.color.rgb=RGBColor(*color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def bullets(doc,text):
    parts=[x.strip(' -•\t') for x in re.split(r'[\n;]+',str(text)) if x.strip()]
    for x in parts or ['TBD']:
        doc.add_paragraph(x,style='List Bullet')

def add_table(doc,rows,headers=("Dimension","Summary")):
    t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers): set_cell_text(t.rows[0].cells[i],h,True,(255,255,255)); shade(t.rows[0].cells[i],'17365D')
    for a,b in rows:
        cells=t.add_row().cells; set_cell_text(cells[0],a,True,(23,54,93)); set_cell_text(cells[1],b)
    return t

def build(d,out):
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
    styles=doc.styles
    styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5)
    for nm,size,col in [('Title',26,(23,54,93)),('Heading 1',17,(47,85,151)),('Heading 2',13,(23,54,93)),('Heading 3',11,(47,85,151))]:
        styles[nm].font.name='Aptos Display'; styles[nm].font.size=Pt(size); styles[nm].font.color.rgb=RGBColor(*col)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('BUSINESS INITIATIVE'); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor(47,85,151)
    title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.add_run(d['initiative_name'])
    sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=sub.add_run(d['one_sentence_promise']); rr.italic=True; rr.font.size=Pt(12)
    add_table(doc,[("Portfolio",d['portfolio']),("Status",d['status']),("Decision requested",d['decision_requested']),("Executive audience",d['executive_audience'])])

    doc.add_heading('1. Executive Snapshot',1)
    add_table(doc,[("Initiative",d['initiative_name']),("Customer problem",d['customer_problem']),("Customer outcome",d['customer_outcome']),("Behavior change",d['behavior_change']),("Value proposition",d['value_proposition']),("Growth mechanism",d['growth_mechanism']),("Why us",d['right_to_win']),("Portfolio role",d['portfolio_leverage']),("Recommendation",d['recommendation']),("Current availability",d['availability_now'])])
    doc.add_heading('Expected outcomes',2); bullets(doc,d['customer_outcome']+'; '+d['north_star'])
    doc.add_heading('Top risks and mitigations',2); bullets(doc,d['risks'])

    doc.add_heading('2. Business Initiative Definition',1)
    for h,k in [('Customer problem','customer_problem'),('Current workflow and evidence','current_workflow'),('Initiative','solution'),('Customer value','customer_outcome'),('How it grows the business','growth_mechanism'),('Why us / right to win','right_to_win'),('Portfolio leverage','portfolio_leverage'),('Scope and boundaries','out_of_scope')]: doc.add_heading(h,2); doc.add_paragraph(d[k])

    doc.add_heading('3. Delivery Roadmap',1)
    add_table(doc,[("NOW",d['availability_now']),("NEAR TERM",d['near_term']),("MEDIUM TERM",d['medium_term']),("BEYOND",d['beyond'])],("Horizon","Capabilities, customer value, and release intent"))
    doc.add_heading('Dependencies and phase gates',2); bullets(doc,d['dependencies']+'; '+d['phase_gates'])

    doc.add_heading('4. Amazon-Style Press Release',1)
    doc.add_heading(f"{d['portfolio']} introduces {d['initiative_name']}",2)
    doc.add_paragraph(f"{d['portfolio']} announces {d['initiative_name']}, an initiative designed to {d['one_sentence_promise'].rstrip('.')}.")
    doc.add_paragraph(f"Customers currently face this problem: {d['customer_problem']} The proposed experience is: {d['solution']}")
    doc.add_paragraph(f"The intended outcome is {d['customer_outcome']} The initiative will be governed by explicit evidence and approval gates.")

    doc.add_heading('5. Amazon FAQ',1)
    faqs=[('What customer problem does this solve?',d['customer_problem']),('What is the proposed solution?',d['solution']),('What does success mean?',d['north_star']),('What is available now?',d['availability_now']),('How does it grow the business?',d['growth_mechanism']),('What must be proven before scaling?',d['phase_gates']),('What is out of scope?',d['out_of_scope']),('How will governance work?',d['governance'])]
    for i,(q,a) in enumerate(faqs,1): doc.add_heading(f'{i}. {q}',3); doc.add_paragraph(a)

    doc.add_heading('6. 1-Page Decision Brief',1)
    add_table(doc,[("Problem",d['customer_problem']),("Audience",d['primary_customer']),("Customer outcome",d['customer_outcome']),("Behavior change",d['behavior_change']),("Growth thesis",d['growth_mechanism']),("Right to win",d['right_to_win']),("Scope",d['solution']),("Stage gates",d['phase_gates']),("North star",d['north_star']),("Recommendation",d['recommendation'])])

    doc.add_heading('7. 6-Pager Narrative',1)
    for h,k in [('Context and strategic fit','portfolio_leverage'),('Customer evidence and validation','problem_evidence'),('Opportunity and customer','primary_customer'),('Proposed approach','solution'),('Business model implications','pricing_hypothesis'),('Operating model','operating_model'),('Milestones and gates','phase_gates'),('Validation priorities','validate_next')]: doc.add_heading(h,2); doc.add_paragraph(d[k])

    doc.add_heading('8. Value Proposition Canvas',1)
    add_table(doc,[("Jobs",d['current_workflow']),("Pains",d['customer_problem']),("Gains",d['customer_outcome']),("Products / services",d['solution']),("Pain relievers",d['value_proposition']),("Fit assessment",d['problem_evidence'])])

    doc.add_heading('9. Commercial Thesis and Ideal Customer Profile',1); doc.add_paragraph(d['growth_mechanism']); doc.add_heading('ICP',2); doc.add_paragraph(d['primary_customer']); doc.add_heading('Disqualifiers',2); bullets(doc,d['out_of_scope'])
    doc.add_heading('10. Buying Committee and Personas',1); doc.add_paragraph(d['buying_committee'])
    doc.add_heading('11. Customer Economic Value',1); doc.add_paragraph(d['customer_outcome']); doc.add_paragraph('Quantify benefits using customer-observed baselines. Do not monetize unsupported risk reduction.')
    doc.add_heading('12. Pricing and Packaging',1); doc.add_paragraph(d['pricing_hypothesis'])
    doc.add_heading('13. Business Impact Model',1); doc.add_paragraph(d['financial_scenarios']); doc.add_paragraph('All values in this section are planning scenarios unless independently validated and approved as a forecast.')
    doc.add_heading('14. Business and Customer Success Metrics',1); bullets(doc,d.get('north_star','TBD')+'; '+d.get('customer_outcome','TBD'))
    doc.add_heading('15. Initiative Scorecard',1); doc.add_paragraph('Score customer severity, impact, strategic alignment, right to win, portfolio leverage, ICP attractiveness, urgency, willingness to pay, monetization, differentiation, feasibility, time to value, and financial attractiveness using evidence gathered during review.')
    doc.add_heading('16. Assumptions and What to Validate Next',1); bullets(doc,d['validate_next']+'; '+d['problem_evidence']+'; '+d['dependencies'])
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('CONFIDENTIAL - EXECUTIVE BUSINESS INITIATIVE'); r.bold=True; r.font.color.rgb=RGBColor(102,102,102)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('CONFIDENTIAL - BUSINESS INITIATIVE')
    doc.core_properties.title=f"{d['initiative_name']} Business Initiative"; doc.core_properties.subject='Executive business initiative proposal'
    doc.save(out)

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--input'); ap.add_argument('--interactive',action='store_true'); ap.add_argument('--output',required=True); a=ap.parse_args()
    if a.interactive:
        d=ask(); side=Path(a.output).with_suffix('.answers.json'); side.write_text(json.dumps(d,indent=2),encoding='utf-8'); print(f"Saved responses: {side}")
    elif a.input: d=json.loads(Path(a.input).read_text(encoding='utf-8'))
    else: ap.error('Use --interactive or --input')
    defaults={k:v for k,_,v in FIELDS}
    defaults.update({k:str(v) for k,v in d.items() if v is not None and str(v).strip()})
    build(defaults,a.output); print(f"Created: {a.output}")
if __name__=='__main__': main()
