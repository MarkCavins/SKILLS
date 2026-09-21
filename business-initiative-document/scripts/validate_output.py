#!/usr/bin/env python3
import sys,re
from pathlib import Path
from docx import Document
p=Path(sys.argv[1]); failures=[]
if not p.exists() or p.stat().st_size<10000: failures.append('file missing or too small')
else:
 d=Document(p); text='\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
 for h in ['Executive Snapshot','Business Initiative Definition','Delivery Roadmap','Amazon-Style Press Release','Amazon FAQ','1-Page Decision Brief','Business Impact Model','Assumptions and What to Validate Next']:
  if h not in text: failures.append('missing heading: '+h)
 if 'Recommendation' not in text: failures.append('missing recommendation')
 if re.search(r'\{\{.*?\}\}|\[\[.*?\]\]',text): failures.append('unreplaced placeholder')
if failures:
 print('FAIL'); [print('- '+x) for x in failures]; sys.exit(1)
print('PASS: structure, recommendation, and placeholders validated')
